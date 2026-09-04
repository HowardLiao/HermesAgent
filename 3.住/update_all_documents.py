import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

import pptx
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PRGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT_DIR = "/Users/howardliao/legal_strategy_project"

# -------------------------------------------------------------
# 1. GENERATE WORD DOCX
# -------------------------------------------------------------
doc = Document()
for section in doc.sections:
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(22)
    section.right_margin = Mm(22)

COLOR_RUST = RGBColor(168, 50, 18)     # #A83212 鐵鏽
COLOR_ORANGE = RGBColor(222, 88, 20)   # #DE5814 標題橘
COLOR_AMBER = RGBColor(217, 119, 6)    # #D97706 重點琥珀黃
COLOR_CITRUS = RGBColor(234, 88, 12)   # #EA580C 柑橘
COLOR_CHARCOAL = RGBColor(45, 55, 72)  # #2D3748 黑灰色內文
FONT_FAMILY = "PingFang TC"

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_title(text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level==1 else 10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = FONT_FAMILY
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = COLOR_ORANGE
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = COLOR_RUST
    else:
        run.font.size = Pt(11.5)
        run.font.color.rgb = COLOR_CITRUS
    return p

def add_body_p(text_runs):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    for item in text_runs:
        if isinstance(item, str):
            text, is_hl, is_b = item, False, False
        else:
            text = item[0]
            is_hl = item[1] if len(item) > 1 else False
            is_b = item[2] if len(item) > 2 else False
        
        run = p.add_run(text)
        run.font.name = FONT_FAMILY
        run.font.bold = is_b or is_hl
        run.font.size = Pt(10.5)
        if is_hl:
            run.font.color.rgb = COLOR_AMBER
        else:
            run.font.color.rgb = COLOR_CHARCOAL
    return p

# Cover
cover_p = doc.add_paragraph()
cover_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
art_path = os.path.join(OUT_DIR, "pharaoh_cover_art.png")
if os.path.exists(art_path):
    doc.add_picture(art_path, width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

t_p = doc.add_paragraph()
t_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
t_p.paragraph_format.space_before = Pt(12)
t_p.paragraph_format.space_after = Pt(4)
r = t_p.add_run("八旬失智配偶權益保全與離婚訴訟策略報告")
r.font.name = FONT_FAMILY
r.font.size = Pt(21)
r.font.bold = True
r.font.color.rgb = COLOR_ORANGE

s_p = doc.add_paragraph()
s_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
s_p.paragraph_format.space_after = Pt(16)
r = s_p.add_run("阻斷惡意脫產・夫妻剩餘財產清算・監護宣告・追討第三者・權益最大化全指南")
r.font.name = FONT_FAMILY
r.font.size = Pt(11.5)
r.font.bold = True
r.font.color.rgb = COLOR_RUST

div_path = os.path.join(OUT_DIR, "egypt_divider.png")
if os.path.exists(div_path):
    doc.add_picture(div_path, width=Inches(5.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Reporter Box
rep_tbl = doc.add_table(rows=1, cols=1)
rep_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
c = rep_tbl.cell(0, 0)
set_cell_shading(c, "FFFBEB")
set_cell_margins(c, 160, 160, 260, 260)
rp = c.paragraphs[0]
rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r1 = rp.add_run("報告人：Howard Liao Ph.D.(廖倫豪 博士)\n")
r1.font.name = FONT_FAMILY
r1.font.size = Pt(13)
r1.font.bold = True
r1.font.color.rgb = COLOR_ORANGE
r2 = rp.add_run("台灣家事法規與權益實務專案 ｜ 日期：2026年9月")
r2.font.name = FONT_FAMILY
r2.font.size = Pt(10)
r2.font.color.rgb = COLOR_CHARCOAL

doc.add_page_break()

# ----------------- 壹、說明篇 -----------------
add_title("壹、【說明篇】案情實務剖析與戰略核心思維", level=1)
add_body_p([
    ("當事人雙方均約八十歲，母親罹患腦神經退化（失智症），日常生活已達需專人照顧且符合外籍看護聘僱資格；父親在外另有第三者（小三），現欲逼迫母親離婚。從台灣家事法實務與司法審判慣例來看，本案具有極為明確之", False, False),
    ("法律與道德不對稱優勢", True, True),
    ("，只要策略得當，母親的權益受法律高度保護。", False, False)
])
add_title("1. 外遇有過失方「無權逼離」之法律鐵壁", level=2)
add_body_p([
    ("依《民法》第 1052 條第 2 項但書規定：「難以維持婚姻之重大事由應由夫妻之一方負責者，僅他方得請求離婚。」", False, False),
    ("父親外遇為破壞婚姻之唯一過失方，母親罹患失智症為重大疾病而非過失", True, True),
    ("。在法理與司法審判上，法院極度排斥「另一半失智生病就想拋棄丟包」之脫法訴求，若母親不同意，父親單方面訴請判決離婚的勝訴機率微乎其微。", False, False)
])
add_title("2. 核心戰略抉擇：「離」與「不離」的極限利益博弈", level=2)
add_body_p([
    ("• ", False, False),
    ("策略 A：堅決不離（以拖待變・保全最大遺產）", True, True),
    ("：母親保留合法配偶之第一順位法定繼承權。父親百年時，母親先主張「夫妻剩餘財產分一半」，剩下遺產再與子女平分。小三在法律上為完全局外人，一毛遺產皆無權繼承。", False, False)
])
add_body_p([
    ("• ", False, False),
    ("策略 B：反訴離婚（即刻清算・落袋為安）", True, True),
    ("：若發現父親正急速賤賣房產、解約定存將資金洗給小三，應立即發動反訴，全面聲請假扣押凍結其資產，分得一半剩餘財產＋追討小三贈與＋高額贍養費，即刻將數千萬產權與存款過戶至母親名下設立安養信託。", False, False)
])

# ----------------- 貳、程序篇 -----------------
add_title("貳、【程序篇】法院訴訟與財產保全三大法定程序", level=1)
add_body_p([
    ("為防範母親權益受損並阻斷父親脫產，律師必須同步啟動以下三大司法程序：", False, False)
])
add_title("1. 人身與意思能力保全：聲請家事法院「監護宣告」程序", level=2)
add_body_p([
    ("依民法第 14 條規定，向管轄地方法院家事法庭具狀聲請。母親因神經退化已無意思辨識能力，由法院裁定", False, False),
    ("指定子女（您）出任法定監護人兼法定代理人", True, True),
    ("。自此，母親名下財產受法院監護，父親無法誘騙母親簽署離婚協議書或過戶切結書；未來所有法庭調解與訴訟，均由監護人代理出庭捍衛權益。", False, False)
])
add_title("2. 財產保全止血程序：聲請「假扣押」與「暫時處分」", level=2)
add_body_p([
    ("為防止父親在訴訟期間將名下不動產脫產移轉至小三名下，應依家事事件法向法院聲請", False, False),
    ("「假扣押裁定」查封主要不動產與主力銀行定存", True, True),
    ("；同時聲請「暫時處分」，強制命令父親按月支付母親外籍看護月薪與醫療長照生活費（民法第1116條之1夫妻互負第一順位扶養義務）。", False, False)
])
add_title("3. 實體裁判程序：民事家事綜合反擊訴訟", level=2)
add_body_p([
    ("在法院審理程序中，由代理人主張駁回父親無理之離婚起訴；同時發動侵害配偶權損害賠償訴訟、夫妻剩餘財產清算訴訟及贍養費給付裁定。", False, False)
])

# ----------------- 參、辦法篇 -----------------
add_title("參、【辦法篇】爭取權益與賠償金最大化四大實體求償辦法", level=1)

claims_tbl = doc.add_table(rows=5, cols=4)
claims_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["求償辦法項目", "民法條文依據", "求償辦法與實務計算重點", "預估效益與戰略定位"]
for i, h in enumerate(headers):
    c = claims_tbl.cell(0, i)
    set_cell_shading(c, "FEF3C7")
    set_cell_margins(c, 100, 100, 120, 120)
    p = c.paragraphs[0]
    r = p.add_run(h)
    r.font.bold = True
    r.font.color.rgb = COLOR_RUST

c_rows = [
    ("一、夫妻剩餘財產分配請求權", "民法第 1030 條之 1", "清查雙方現存之婚後財產扣除負債後差額平均分配。八十歲長輩累積之房地產與高額定存多屬婚後取得，母親依法享有差額一半（50%）請求權。", "【金額最大宗・佔90%以上】\n動輒數百萬至數千萬元法定產權，為母親晚年生活最堅實保障。"),
    ("二、追討第三者條款（追加計算）", "民法第 1030 條之 3\n民法第 1020 條之 1", "起訴前五年內，父親為減少分配而無償贈與第三者之資產，全數「追加計算，視為父親現存之婚後財產」！一年內無償贈與得直接訴請法院撤銷返還。", "【破解洗錢脫產殺手鐧】\n縱使父親偷偷買房、匯巨款給小三，全數加回視同在案，判男方補償。"),
    ("三、侵害配偶權連帶精神撫慰金", "民法第 184 條\n民法第 195 條第 3 項", "共同提告父親與小三為連帶被告。舉證兩人同居、出遊牽手、親暱對話。審酌八旬原配病中遭背叛之重大痛苦，判命兩人連帶賠償。", "【直擊第三者承擔賠償責任】\n實務判賠金額普遍落於 30 萬 ～ 100 萬元不等。"),
    ("四、離婚損害賠償與終身贍養費", "民法第 1056 條\n民法第 1057 條", "判決離婚非財產損害賠償；加計母親八旬失智、無謀生能力且生活陷入重大困難，判命父親按月支付終身扶養照護費或一次性大額贍養費。", "【終身全額保障尊嚴】\n外籍看護月薪（3.5萬）+醫療長照耗材由父親全額埋單。")
]

for row_i, r_data in enumerate(c_rows, start=1):
    for col_i, text in enumerate(r_data):
        c = claims_tbl.cell(row_i, col_i)
        set_cell_shading(c, "FFFDF7" if row_i % 2 == 0 else "FFFFFF")
        set_cell_margins(c, 90, 90, 110, 110)
        p = c.paragraphs[0]
        r = p.add_run(text)
        r.font.name = FONT_FAMILY
        r.font.size = Pt(9.5)
        if col_i == 0:
            r.font.bold = True
            r.font.color.rgb = COLOR_ORANGE
        elif col_i == 3 and "【" in text:
            r.font.bold = True
            r.font.color.rgb = COLOR_RUST
        else:
            r.font.color.rgb = COLOR_CHARCOAL

# ----------------- 肆、步驟篇 -----------------
add_title("肆、【步驟篇】家屬即刻行動指南與實施六大步驟（SOP）", level=1)

add_title("步驟一：前往戶籍地區公所領取【身心障礙鑑定表】", level=2)
add_body_p([
    ("攜帶證件：母親身分證正本、印章、1吋照片3張、代辦人身分證與印章。先領取空白表格，以便看診當天同步完成醫療鑑定，免除多跑一趟醫院。", False, False)
])

add_title("步驟二：台中榮總神經內科就診與「精準專業診斷書」開立", level=2)
add_body_p([
    ("依律師指示，請醫師在【乙種診斷證明書】中具體載明以下專業醫學與法定術語：", False, False)
])
add_body_p([
    ("1. 醫學專有名稱：載明「阿茲海默型失智症／進行性腦神經退化症」伴隨「嚴重認知功能障礙」。", False, False)
])
add_body_p([
    ("2. 記憶力與定向感：載明「近期與遠期記憶力顯著缺損，無法回憶重大生活與金錢事件」、「人、時、地定向感顯著喪失（Disorientation）」。", False, False)
])
add_body_p([
    ("3. 智力與判斷力：載明「抽象思考、計算能力與判斷力嚴重衰退」。", False, False)
])
add_body_p([
    ("4. 客觀量表指標：務必列出「臨床失智評估量表（CDR）＝ 1分以上」及 MMSE 測驗分數。", False, False)
])
add_body_p([
    ("5. 法定無辨識能力（最關鍵字句）：", True, True),
    ("醫囑註記「病人因神經退化致心智缺陷與精神障礙，", False, False),
    ("致不能為意思表示或受意思表示，已無獨立處理自己事務之能力", True, True),
    ("，日常生活活動能力完全喪失，需專人24小時全天候看護照料」。", False, False)
])
add_body_p([
    ("★ 同步於台中榮總第一醫療大樓一樓身障窗口遞交身障鑑定表，由院方專人送衛生局審核發證。", False, False)
])

add_title("步驟三：向國稅局調閱父母雙方財產總歸戶與所得清單", level=2)
add_body_p([
    ("調閱「全國財產稅總歸戶財產查詢清單」與「近三年綜合所得稅各類所得清單」。全面掌握父親名下所有土地建號、上市櫃股票集保庫存，並由銀行利息所得反推巨額存款帳戶，徹底摸清父親資產底牌。", False, False)
])

add_title("步驟四：向家事法院具狀聲請「監護宣告」", level=2)
add_body_p([
    ("檢附榮總診斷書與戶籍謄本，由子女向法院聲請對母親做監護宣告並由子女擔任監護人。徹底阻斷父親私下誘騙母親簽字辦理協議離婚或過戶產權之任何可能性。", False, False)
])

add_title("步驟五：向法院聲請「假扣押／暫時處分」凍結資產", level=2)
add_body_p([
    ("由律師向法院聲請查封父親主要房產與高額存款帳戶，禁止其變賣、移轉或脫產給小三；同步聲請暫時處分，命父親於訴訟確定前按月匯給外籍看護薪資與生活費。", False, False)
])

add_title("步驟六：實體訴訟反擊、共同求償與分產談判", level=2)
add_body_p([
    ("提告父親與小三連帶賠償侵害配偶權百萬慰撫金；在家事法庭開庭時，依選定之戰略（策略A阻斷離婚保全遺產，或策略B反訴平分剩餘財產＋追討追加小三資產），由監護人代表出庭爭取全勝！", False, False)
])

doc_out_path = os.path.join(OUT_DIR, "legal_strategy_report.docx")
doc.save(doc_out_path)
print("Word Docx rebuilt successfully at:", doc_out_path)

